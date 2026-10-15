import datetime as dt
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from server import excel_reader, excel_writer
from server.app import create_app
from server.storage import Conflict, Store


def make_book(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Label"
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="2F75B5")
    ws["B1"] = 10
    ws["B2"] = "=B1*2"
    ws["C1"] = 0.25
    ws["C1"].number_format = "0.0%"
    ws["D1"] = dt.datetime(2026, 9, 27)
    ws["D1"].number_format = "dd-mmm-yyyy"
    ws["E1"] = "007"
    ws.merge_cells("A3:C4")
    ws["A3"] = "Merged"
    wb.create_sheet("Other")["A1"] = "x"
    wb.save(path)
    return path


@pytest.fixture()
def root(tmp_path):
    d = tmp_path / "Dashboard" / "M" / "MS" / "V"
    d.mkdir(parents=True)
    make_book(d / "Doc.xlsx")
    return tmp_path / "Dashboard"


def test_format_number():
    assert excel_reader.format_number(1234.5, "#,##0.00") == "1,234.50"
    assert excel_reader.format_number(0.125, "0.0%") == "12.5%"
    assert excel_reader.format_number(3, "General") == "3"
    assert excel_reader.format_number(2.0, "General") == "2"
    assert excel_reader.format_number(0.1 + 0.2, "General") == "0.3"


def test_format_date():
    assert excel_reader.format_date(dt.datetime(2026, 9, 7, 14, 5), "dd-mmm-yyyy") == "07-Sep-2026"
    assert excel_reader.format_date(dt.datetime(2026, 9, 7, 14, 5), "d/m/yy h:mm") == "7/9/26 14:05"


def test_theme_tint():
    assert excel_reader.apply_tint("4472C4", 0) == "4472C4"
    assert excel_reader.apply_tint("000000", 0.5) == "808080"


def test_coerce_matches_original_type():
    assert excel_writer.coerce("12", 10, "General") == 12
    assert excel_writer.coerce("12.5%", 0.1, "0%") == 0.125
    assert excel_writer.coerce("4.30", "4.30", "General") == "4.30"  # text stays text
    assert excel_writer.coerce("007", None, "General") == "007"
    assert excel_writer.coerce("", 5, "General") is None
    assert isinstance(excel_writer.coerce("=A1*2", 1, "General"), excel_writer.Formula)


def test_reader_model(root):
    m = excel_reader.read_workbook(str(root / "M/MS/V/Doc.xlsx"))
    sh = m["sheets"][0]
    assert sh["name"] == "Data"
    assert sh["cells"]["1,3"]["v"] == "25.0%"
    assert sh["cells"]["1,3"]["e"] == "25%"
    assert sh["cells"]["1,4"]["v"] == "27-Sep-2026"
    assert sh["cells"]["2,2"]["f"] == "=B1*2"
    assert {"r1": 3, "c1": 1, "r2": 4, "c2": 3} in sh["merges"]
    style = sh["styles"][sh["cells"]["1,1"]["s"]]
    assert "background:#2F75B5" in style and "font-weight:700" in style


def test_writer_preserves_package(root):
    src = (root / "M/MS/V/Doc.xlsx").read_bytes()
    out = excel_writer.apply_edits(src, {"Data": {(1, 2): 42, (6, 6): "new cell", (1, 1): "Renamed"}})
    wb = load_workbook(io.BytesIO(out))
    ws = wb["Data"]
    assert ws["B1"].value == 42
    assert ws["F6"].value == "new cell"
    assert ws["A1"].value == "Renamed"
    assert ws["A1"].font.b and ws["A1"].fill.fgColor.rgb.endswith("2F75B5")  # style kept
    assert ws["B2"].value == "=B1*2"
    assert wb["Other"]["A1"].value == "x"
    names_in = {i.filename for i in zipfile.ZipFile(io.BytesIO(src)).infolist()}
    names_out = {i.filename for i in zipfile.ZipFile(io.BytesIO(out)).infolist()}
    assert names_in - {"xl/calcChain.xml"} == names_out - {"xl/calcChain.xml"}


def test_save_history_and_conflict(root):
    store = Store(str(root))
    rel = "M/MS/V/Doc.xlsx"
    doc = store.load(rel)
    assert doc["rev"] == 0
    res = store.save(rel, {"Data": {"1,2": "11", "1,5": "007"}}, "Alice", "bump", 0)
    assert res["saved"] and res["rev"] == 1
    assert [c["cell"] for c in res["changes"]] == ["B1"]  # E1 unchanged -> not logged
    with pytest.raises(Conflict):
        store.save(rel, {"Data": {"1,2": "12"}}, "Bob", "", 0)
    # external edit is detected and diffed
    f = root / rel
    wb = load_workbook(f)
    wb["Data"]["B1"] = 99
    wb.save(f)
    revs = store.revisions(rel)
    assert revs[-1]["kind"] == "external"
    assert revs[-1]["changes"][0]["cell"] == "B1" and revs[-1]["changes"][0]["new"] == "99"
    store.restore(rel, 1, "Alice")
    assert load_workbook(f)["Data"]["B1"].value == 11


def test_path_traversal_blocked(root):
    store = Store(str(root))
    with pytest.raises(PermissionError):
        store.resolve("../secret.xlsx")
    with pytest.raises(PermissionError):
        store.resolve(".dochub/activity.jsonl")


def test_api_roundtrip(root):
    app = create_app(str(root))
    c = app.test_client()
    tree = c.get("/api/tree").get_json()
    assert tree["models"][0]["milestones"][0]["variants"][0]["documents"][0]["name"] == "Doc"
    doc = c.get("/api/document?path=M/MS/V/Doc.xlsx").get_json()
    r = c.post("/api/save", json={"path": "M/MS/V/Doc.xlsx", "baseRev": doc["rev"], "user": "T",
                                   "note": "n", "edits": {"Data": {"1,1": "Hello"}}})
    assert r.status_code == 200 and r.get_json()["rev"] == 1
    r = c.post("/api/export", json={"path": "M/MS/V/Doc.xlsx", "edits": {"Data": {"1,1": "Draft"}}})
    assert load_workbook(io.BytesIO(r.data))["Data"]["A1"].value == "Draft"
    assert c.get("/api/stats").get_json()["revisions"] == 1
    assert c.get("/api/document?path=../x.xlsx").status_code == 403


def _sheet_xml(cells_xml):
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="{excel_writer.NS}">'
            f'<sheetData>{cells_xml}</sheetData></worksheet>').encode()


def test_shared_formula_master_edit_expands_group():
    xml = _sheet_xml('<row r="1"><c r="A1"><v>1</v></c><c r="B1"><f t="shared" ref="B1:B3" si="0">A1*2</f><v>2</v></c></row>'
                     '<row r="2"><c r="A2"><v>2</v></c><c r="B2"><f t="shared" si="0"/><v>4</v></c></row>'
                     '<row r="3"><c r="A3"><v>3</v></c><c r="B3"><f t="shared" si="0"/><v>6</v></c></row>')
    out = excel_writer._patch_sheet(xml, {(1, 2): 99}).decode()
    assert "<f>A2*2</f>" in out and "<f>A3*2</f>" in out and 'si="0"' not in out


def test_array_formula_member_edit_rejected():
    xml = _sheet_xml('<row r="1"><c r="A1"><f t="array" ref="A1:A2">B1:B2*2</f><v>1</v></c></row>'
                     '<row r="2"><c r="A2"><v>2</v></c></row>')
    with pytest.raises(excel_writer.EditError):
        excel_writer._patch_sheet(xml, {(2, 1): 5})


def test_api_legacy_and_encrypted_files_explain_themselves(root):
    v = root / "M/MS/V"
    (v / "Old.xls").write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + bytes(504))
    (v / "Secret.xlsx").write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + "EncryptedPackage".encode("utf-16-le") + bytes(64))
    c = create_app(str(root)).test_client()
    docs = c.get("/api/tree").get_json()["models"][0]["milestones"][0]["variants"][0]["documents"]
    assert {d["file"]: d["supported"] for d in docs} == {"Doc.xlsx": True, "Old.xls": False, "Secret.xlsx": True}
    r = c.get("/api/document?path=M/MS/V/Old.xls")
    assert r.status_code == 422 and r.get_json()["code"] == "legacy"
    r = c.get("/api/document?path=M/MS/V/Secret.xlsx")
    assert r.status_code == 422 and r.get_json()["code"] == "encrypted"


def test_documents_in_subfolders_and_milestones(root):
    (root / "M/MS/V/Sub").mkdir()
    make_book(root / "M/MS/V/Sub/Deep.xlsx")
    make_book(root / "M/MS/Loose.xlsx")
    t = create_app(str(root)).test_client().get("/api/tree").get_json()
    ms = t["models"][0]["milestones"][0]
    assert [d["name"] for d in ms["documents"]] == ["Loose"]
    deep = [d for d in ms["variants"][0]["documents"] if d["name"] == "Deep"][0]
    assert deep["folder"] == "Sub" and deep["path"] == "M/MS/V/Sub/Deep.xlsx"


def test_edited_by_uses_excel_author_then_dochub_user(root):
    f = root / "M/MS/V/Doc.xlsx"
    wb = load_workbook(f)
    wb.properties.lastModifiedBy = "Priya (Excel)"
    wb.save(f)
    c = create_app(str(root)).test_client()
    doc = c.get("/api/tree").get_json()["models"][0]["milestones"][0]["variants"][0]["documents"][0]
    assert doc["editedBy"] == "Priya (Excel)" and doc["editedAt"]
    c.post("/api/save", json={"path": "M/MS/V/Doc.xlsx", "baseRev": None, "user": "Alice", "note": "",
                              "edits": {"Data": {"1,1": "Changed"}}})
    doc = c.get("/api/tree").get_json()["models"][0]["milestones"][0]["variants"][0]["documents"][0]
    assert doc["editedBy"] == "Alice"
    assert load_workbook(f).properties.lastModifiedBy == "Alice"  # Excel's File > Info shows it too


def test_column_width_matches_excel():
    assert excel_reader.col_width_px(9.140625) == 64   # Excel default column
    assert excel_reader.col_width_px(30) == 210


def test_runs_when_compiled_addons_are_blocked(tmp_path):
    """Corporate Application Control policies can block DLLs such as lxml / Pillow."""
    import subprocess
    import sys
    from pathlib import Path
    code = r'''
import importlib.abc, sys
class Block(importlib.abc.MetaPathFinder):
    def find_spec(self, name, path=None, target=None):
        if name.split(".")[0] in ("lxml", "PIL"):
            raise ImportError("DLL load failed: An Application Control policy has blocked this file.")
sys.meta_path.insert(0, Block())
import run  # noqa: F401  (imports the whole server)
from server import excel_reader, excel_writer
from openpyxl import Workbook, load_workbook
wb = Workbook(); wb.active["A1"] = "x"; wb.save(sys.argv[1])
excel_reader.read_workbook(sys.argv[1])
out = excel_writer.apply_edits(open(sys.argv[1], "rb").read(), {"Sheet": {(1, 1): "y"}})
open(sys.argv[1], "wb").write(out)
assert load_workbook(sys.argv[1]).active["A1"].value == "y"
assert not [m for m in sys.modules if m.split(".")[0] in ("lxml", "PIL")]
print("ok")
'''
    root = Path(__file__).resolve().parent.parent
    res = subprocess.run([sys.executable, "-c", code, str(tmp_path / "t.xlsx")], cwd=root,
                         capture_output=True, text=True)
    assert res.stdout.strip() == "ok", res.stderr


def test_text_writer_prefixed_namespace_new_rows_and_implicit_cells():
    ns = excel_writer.NS
    xml = (f'<?xml version="1.0"?><x:worksheet xmlns:x="{ns}" xmlns:mc="m" mc:Ignorable="x14ac">'
           '<x:sheetData><x:row r="2" spans="1:2"><x:c><x:v>1</x:v></x:c><x:c t="s"><x:v>0</x:v></x:c></x:row>'
           '<x:row r="9"><x:c r="A9"><x:v>9</x:v></x:c></x:row></x:sheetData></x:worksheet>').encode()
    out = excel_writer._patch_sheet(xml, {(2, 3): "new", (5, 1): 5, (12, 2): True}).decode()
    assert 'mc:Ignorable="x14ac"' in out and "<x:worksheet" in out          # untouched bytes kept
    assert '<x:row r="2"><x:c r="A2"><x:v>1</x:v></x:c><x:c r="B2" t="s">' in out  # implicit cells addressed
    assert '<x:c r="C2" t="inlineStr"><x:is><x:t xml:space="preserve">new</x:t>' in out
    assert out.index('<x:row r="5">') < out.index('<x:row r="9">') < out.index('<x:row r="12">')
    assert '<x:c r="B12" t="b"><x:v>1</x:v></x:c>' in out
    import xml.etree.ElementTree as ET
    ET.fromstring(out.encode())
